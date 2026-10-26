"""FR-018 import & export."""
import csv
import io
import json
import os
import re
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.core import serializers as dj_serializers
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from django.utils.text import slugify

from core import audit
from core.models import SystemSetting
from core.normalize import normalize_phone
from talent import services as tsvc
from talent.models import ContactNote, FeaturedVideo, Talent, TalentLanguage, Term

from .models import ExportJob, ImportBatch, ImportRow

# ---------------------------------------------------------------------------
# Import field catalogue
# ---------------------------------------------------------------------------
IMPORT_FIELDS = {
    "full_name": ("Full name", ["name", "full name", "talent name", "artist name"]),
    "screen_name": ("Screen name", ["screen name", "stage name"]),
    "phone": ("Primary phone", ["phone", "mobile", "phone number", "contact", "contact number", "mobile number", "whatsapp number"]),
    "alt_phone": ("Alternate phone", ["alternate phone", "alt phone", "phone 2", "secondary phone", "other phone"]),
    "whatsapp_phone": ("WhatsApp (if different)", ["whatsapp"]),
    "email": ("Email", ["email", "e-mail", "mail"]),
    "instagram": ("Instagram", ["instagram", "insta", "ig", "instagram link", "instagram url"]),
    "city": ("City / current location", ["city", "location", "place", "current location", "town"]),
    "district": ("District", ["district"]),
    "state": ("State", ["state"]),
    "gender": ("Gender", ["gender", "sex"]),
    "date_of_birth": ("Date of birth", ["dob", "date of birth", "birth date", "birthday"]),
    "approx_age": ("Age", ["age"]),
    "playing_age_min": ("Playing age from", ["playing age min", "playing age from", "min age"]),
    "playing_age_max": ("Playing age to", ["playing age max", "playing age to", "max age"]),
    "playing_age": ("Playing age range (e.g. 25-32)", ["playing age", "age range"]),
    "height_cm": ("Height (cm)", ["height", "height cm"]),
    "categories": ("Categories (comma separated)", ["category", "categories", "type"]),
    "tags": ("Tags (comma separated)", ["tags", "tag", "looks"]),
    "skills": ("Skills (comma separated)", ["skills", "skill", "talents"]),
    "languages": ("Languages (comma separated)", ["languages", "language", "languages known"]),
    "experience_level": ("Experience level", ["experience", "experience level"]),
    "expected_fee": ("Expected fee", ["fee", "expected fee", "rate", "remuneration", "charges"]),
    "fee_type": ("Fee type", ["fee type", "rate type", "fee basis"]),
    "travel_willingness": ("Travel willingness", ["travel", "willing to travel"]),
    "appearance_notes": ("Appearance notes", ["appearance", "look", "appearance notes"]),
    "experience_summary": ("Experience / credits", ["credits", "previous work", "experience summary", "portfolio"]),
    "internal_notes": ("Internal notes", ["notes", "remarks", "comments"]),
    "photo_filename": ("Photo filename (inside ZIP)", ["photo", "photo filename", "image", "photo file"]),
    "video_1": ("Featured video 1 URL", ["video", "video 1", "video link", "showreel"]),
    "video_2": ("Featured video 2 URL", ["video 2"]),
    "video_3": ("Featured video 3 URL", ["video 3"]),
}

GENDER_MAP = {"m": "male", "male": "male", "man": "male", "boy": "male", "f": "female", "female": "female", "woman": "female", "girl": "female",
              "nb": "non_binary", "non-binary": "non_binary", "non binary": "non_binary", "other": "other"}
FEE_TYPE_MAP = {"day": "per_day", "per day": "per_day", "daily": "per_day", "shoot": "per_shoot", "per shoot": "per_shoot",
                "project": "per_project", "per project": "per_project", "hour": "per_hour", "per hour": "per_hour", "hourly": "per_hour"}
TRAVEL_MAP = {"local": "local", "no": "local", "district": "district", "state": "state", "kerala": "state", "yes": "national",
              "india": "national", "national": "national", "anywhere": "national", "international": "international", "abroad": "international"}
EXPERIENCE_MAP = {k: k for k, _ in Talent.EXPERIENCE_CHOICES}
EXPERIENCE_MAP.update({"fresher": "fresher", "new": "fresher", "beginner": "beginner", "intermediate": "intermediate", "experienced": "experienced", "pro": "professional", "professional": "professional"})


def _norm_header(h):
    return re.sub(r"[^a-z0-9 ]", " ", str(h or "").lower()).strip()


def auto_mapping(headers):
    mapping = {}
    used = set()
    for h in headers:
        nh = re.sub(r"\s+", " ", _norm_header(h))
        for key, (label, aliases) in IMPORT_FIELDS.items():
            if key in used:
                continue
            nl = re.sub(r"\s+", " ", _norm_header(label))
            if nh == key.replace("_", " ") or nh == nl or nh in aliases:
                mapping[h] = key
                used.add(key)
                break
    return mapping


# ---------------------------------------------------------------------------
# Reading files
# ---------------------------------------------------------------------------
class ImportFileError(ValueError):
    pass


def read_table(fileobj, name):
    ext = os.path.splitext(name.lower())[1]
    if ext == ".csv":
        raw = fileobj.read()
        for enc in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportFileError("We couldn't read the text in this CSV. Save it as UTF-8 and try again.")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl

        try:
            wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
        except Exception:
            raise ImportFileError("This Excel file couldn't be opened. Make sure it's a valid .xlsx file.")
        ws = wb.worksheets[0]
        rows = [["" if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        raise ImportFileError("Upload a .csv or .xlsx file.")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise ImportFileError("The file needs a header row and at least one data row.")
    headers = [str(h).strip() or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    seen = {}
    for i, h in enumerate(headers):  # make headers unique
        if h in seen:
            seen[h] += 1
            headers[i] = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
    data = []
    for r in rows[1:]:
        cells = list(r) + [""] * (len(headers) - len(r))
        data.append({headers[i]: _cell(cells[i]) for i in range(len(headers))})
    if len(data) > 5000:
        raise ImportFileError("Imports are limited to 5,000 rows at a time. Split the file and try again.")
    return headers, data


def _cell(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def _parse_date(s):
    s = (s or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"“{s}” isn't a date we recognise (use DD/MM/YYYY).")


def _parse_int(s, label, lo=0, hi=300):
    s = (s or "").strip()
    if not s:
        return None
    m = re.search(r"\d+", s)
    if not m:
        raise ValueError(f"{label} “{s}” isn't a number.")
    v = int(m.group())
    if not lo <= v <= hi:
        raise ValueError(f"{label} {v} looks wrong.")
    return v


def _parse_money(s):
    s = (s or "").strip().lower().replace(",", "").replace("₹", "").replace("rs", "").replace("inr", "").strip()
    if not s:
        return None
    mult = 1
    if s.endswith("k"):
        mult, s = 1000, s[:-1]
    elif s.endswith("l") or s.endswith("lakh"):
        mult, s = 100000, s.rstrip("lakh").rstrip("l")
    try:
        return str((Decimal(s.strip()) * mult).quantize(Decimal("0.01")))
    except InvalidOperation:
        raise ValueError(f"Fee “{s}” isn't a number.")


def _split(s):
    return [p.strip() for p in re.split(r"[,;/|]", s or "") if p.strip()]


def mapped_values(row, mapping):
    out = {}
    for header, key in mapping.items():
        if key and header in row and str(row[header]).strip():
            out[key] = str(row[header]).strip()
    return out


def validate_row(values, photo_names, seen_phones):
    errors, warnings, data = [], [], {}
    data["full_name"] = values.get("full_name", "")
    if not data["full_name"]:
        errors.append("Full name is missing.")
    data["phone"] = values.get("phone", "")
    cc = SystemSetting.get().default_country_code
    norm = normalize_phone(data["phone"], cc)
    if not norm:
        errors.append("Primary phone is missing.")
    elif len(norm) < 10:
        errors.append(f"Phone “{data['phone']}” looks incomplete.")
    elif norm in seen_phones:
        errors.append(f"Same phone as row {seen_phones[norm]} in this file.")
    data["city"] = values.get("city", "")
    if not data["city"]:
        errors.append("City / current location is missing.")
    cats = _split(values.get("categories", ""))
    data["categories"] = cats
    if not cats:
        errors.append("At least one category is required.")
    for k in ("screen_name", "alt_phone", "whatsapp_phone", "email", "instagram", "district", "state", "appearance_notes", "experience_summary", "internal_notes"):
        if values.get(k):
            data[k] = values[k]
    for k in ("tags", "skills", "languages"):
        if values.get(k):
            data[k] = _split(values[k])
    if values.get("gender"):
        g = GENDER_MAP.get(values["gender"].strip().lower())
        if g:
            data["gender"] = g
        else:
            warnings.append(f"Gender “{values['gender']}” not recognised — left blank.")
    try:
        dob = _parse_date(values.get("date_of_birth"))
        if dob:
            if dob > timezone.localdate():
                raise ValueError("Date of birth is in the future.")
            data["date_of_birth"] = dob.isoformat()
    except ValueError as e:
        errors.append(str(e))
    for k, label, hi in (("approx_age", "Age", 120), ("playing_age_min", "Playing age from", 120), ("playing_age_max", "Playing age to", 120), ("height_cm", "Height", 260)):
        try:
            v = _parse_int(values.get(k), label, 0, hi)
            if v is not None:
                data[k] = v
        except ValueError as e:
            errors.append(str(e))
    if values.get("playing_age") and "playing_age_min" not in data:
        nums = re.findall(r"\d+", values["playing_age"])
        if len(nums) >= 2:
            data["playing_age_min"], data["playing_age_max"] = int(nums[0]), int(nums[1])
        elif nums:
            data["playing_age_min"] = data["playing_age_max"] = int(nums[0])
    if data.get("playing_age_min") and data.get("playing_age_max") and data["playing_age_min"] > data["playing_age_max"]:
        errors.append("Playing age 'from' is greater than 'to'.")
    if values.get("height_cm") and data.get("height_cm") and data["height_cm"] < 100 and "." in values["height_cm"]:
        # "5.6" feet → cm
        try:
            ft, inch = values["height_cm"].split(".")
            data["height_cm"] = round(int(ft) * 30.48 + int(inch) * 2.54)
            warnings.append(f"Height converted from feet to {data['height_cm']} cm.")
        except ValueError:
            pass
    try:
        fee = _parse_money(values.get("expected_fee"))
        if fee is not None:
            data["expected_fee"] = fee
    except ValueError as e:
        errors.append(str(e))
    for k, table in (("fee_type", FEE_TYPE_MAP), ("travel_willingness", TRAVEL_MAP), ("experience_level", EXPERIENCE_MAP)):
        if values.get(k):
            v = table.get(values[k].strip().lower())
            if v:
                data[k] = v
            else:
                warnings.append(f"{k.replace('_', ' ').title()} “{values[k]}” not recognised — left blank.")
    videos = [values.get(f"video_{i}") for i in (1, 2, 3) if values.get(f"video_{i}")]
    good_videos = []
    for v in videos:
        if re.match(r"^https?://", v):
            good_videos.append(v)
        else:
            warnings.append(f"Video “{v[:40]}” isn't a web link — skipped.")
    if good_videos:
        data["videos"] = good_videos
    photo = values.get("photo_filename", "")
    if photo:
        if os.path.basename(photo).lower() in photo_names:
            data["photo_filename"] = photo_names[os.path.basename(photo).lower()]
        else:
            warnings.append(f"Photo “{photo}” wasn't found in the ZIP.")
    if norm and norm not in seen_phones:
        seen_phones[norm] = None  # filled by caller
    return data, errors, warnings


def _zip_names(batch):
    if not batch.photos_zip:
        return {}
    try:
        with batch.photos_zip.open("rb") as fh, zipfile.ZipFile(fh) as z:
            return {
                os.path.basename(n).lower(): n
                for n in z.namelist()
                if not n.endswith("/") and os.path.splitext(n.lower())[1] in (".jpg", ".jpeg", ".png", ".webp") and not os.path.basename(n).startswith(".")
            }
    except zipfile.BadZipFile:
        return {}


def validate_batch(batch, mapping):
    """Store per-row preview: data, errors, warnings, duplicates, default action."""
    batch.mapping = mapping
    photo_names = _zip_names(batch)
    seen = {}
    counts = {"ok": 0, "errors": 0, "duplicates": 0, "staging": 0, "warnings": 0}
    cc = SystemSetting.get().default_country_code
    rows = list(batch.rows.all())
    for row in rows:
        values = mapped_values(row.raw, mapping)
        data, errors, warnings = validate_row(values, photo_names, seen)
        norm = normalize_phone(data.get("phone"), cc)
        if norm and seen.get(norm) is None:
            seen[norm] = row.row_number
        dups = []
        if not errors:
            matches = tsvc.find_duplicates(
                full_name=data.get("full_name", ""), phone=data.get("phone", ""), alt_phone=data.get("alt_phone", ""),
                instagram=data.get("instagram", ""), city=data.get("city", ""), age=data.get("approx_age"), limit=5,
            )
            dups = [{"id": m["talent"].id, "cast_id": m["talent"].cast_id, "full_name": m["talent"].full_name, "city": m["talent"].city,
                     "strength": m["strength"], "reasons": m["reasons"]} for m in matches]
        row.data, row.errors, row.warnings, row.duplicates = data, errors, warnings, dups
        if errors:
            row.action = "skip"
            counts["errors"] += 1
        elif any(d["strength"] in ("exact", "strong") for d in dups):
            row.action = "skip"  # never silently create — staff must choose
            counts["duplicates"] += 1
        elif "photo_filename" not in data:
            row.action = "staging"
            counts["staging"] += 1
        else:
            row.action = "create"
            counts["ok"] += 1
        if dups and not any(d["strength"] in ("exact", "strong") for d in dups):
            counts["duplicates"] += 1
        if warnings:
            counts["warnings"] += 1
        row.duplicate_acknowledged = False
    ImportRow.objects.bulk_update(rows, ["data", "errors", "warnings", "duplicates", "action", "duplicate_acknowledged"])
    batch.status = "mapped"
    batch.summary = counts
    batch.save(update_fields=["mapping", "status", "summary"])
    return counts


def _terms(kind, names, user, created_log):
    out = []
    for n in names:
        slug = slugify(n)[:90]
        if not slug:
            continue
        t = Term.objects.filter(kind=kind, slug=slug).first()
        if t and t.merged_into_id:
            t = t.merged_into
        if not t:
            t = Term.objects.create(kind=kind, name=n.strip().title() if n.islower() else n.strip(), created_by=user, updated_by=user)
            created_log.append(f"{kind}: {t.name}")
        out.append(t)
    return out


@transaction.atomic
def _create_from_row(row, batch, user, zf, created_terms):
    d = dict(row.data)
    staging = row.action == "staging" or "photo_filename" not in d
    simple = {k: d[k] for k in (
        "full_name", "screen_name", "phone", "alt_phone", "whatsapp_phone", "email", "instagram", "city", "district", "state",
        "gender", "approx_age", "playing_age_min", "playing_age_max", "height_cm", "expected_fee", "fee_type", "travel_willingness",
        "experience_level", "appearance_notes", "experience_summary", "internal_notes",
    ) if k in d}
    if "date_of_birth" in d:
        simple["date_of_birth"] = date.fromisoformat(d["date_of_birth"])
    if "expected_fee" in simple:
        simple["expected_fee"] = Decimal(simple["expected_fee"])
        simple["rates_reviewed_at"] = timezone.now()
    t = Talent(**simple, created_by=user, updated_by=user, status=Talent.STATUS_STAGING if staging else Talent.STATUS_NEEDS_DETAILS)
    t.save()
    t.categories.set(_terms(Term.CATEGORY, d.get("categories", []), user, created_terms))
    t.tags.set(_terms(Term.TAG, d.get("tags", []), user, created_terms))
    t.skills.set(_terms(Term.SKILL, d.get("skills", []), user, created_terms))
    for lang in _terms(Term.LANGUAGE, d.get("languages", []), user, created_terms):
        TalentLanguage.objects.get_or_create(talent=t, language=lang)
    for i, url in enumerate(d.get("videos", [])[:3], start=1):
        FeaturedVideo.objects.create(
            talent=t, position=i, title=f"Video {i}", url=url, platform=tsvc.detect_platform(url), thumbnail_url=tsvc.video_thumbnail(url),
            created_by=user, updated_by=user,
        )
    if not staging and zf is not None:
        raw = zf.read(d["photo_filename"])
        upload = ContentFile(raw, name=os.path.basename(d["photo_filename"]))
        tsvc.add_photo(t, upload, user=user, make_primary=True)
    if d.get("internal_notes"):
        ContactNote.objects.create(talent=t, kind="note", content=f"Imported note: {d['internal_notes']}", staff=user, created_by=user)
    tsvc.refresh_completeness(t)
    return t


def run_import(batch, user, request=None):
    batch.status = "importing"
    batch.save(update_fields=["status"])
    created = staged = skipped = failed = 0
    created_terms = []
    zf = None
    if batch.photos_zip:
        try:
            zf = zipfile.ZipFile(batch.photos_zip.open("rb"))
        except zipfile.BadZipFile:
            zf = None
    try:
        for row in batch.rows.all():
            if row.action == "skip" or row.errors:
                row.result, row.result_message = "skipped", ("; ".join(row.errors) if row.errors else "Skipped by staff")
                skipped += 1
                row.save(update_fields=["result", "result_message"])
                continue
            strong = [x for x in row.duplicates if x["strength"] in ("exact", "strong")]
            if (strong or row.duplicates) and not row.duplicate_acknowledged:
                row.result, row.result_message = "skipped", "Possible duplicate not acknowledged"
                skipped += 1
                row.save(update_fields=["result", "result_message"])
                continue
            try:
                t = _create_from_row(row, batch, user, zf, created_terms)
                row.talent = t
                if t.status == Talent.STATUS_STAGING:
                    row.result, row.result_message = "staged", f"{t.cast_id} in import staging (photo needed)"
                    staged += 1
                else:
                    row.result, row.result_message = "created", t.cast_id
                    created += 1
            except tsvc.PhotoError as e:
                row.result, row.result_message = "failed", f"Photo problem: {e}"
                failed += 1
            except Exception as e:  # keep going; report per row
                row.result, row.result_message = "failed", f"Couldn't save this row: {e}"[:255]
                failed += 1
            row.save(update_fields=["result", "result_message", "talent"])
    finally:
        if zf:
            zf.close()
    batch.status = "completed"
    batch.completed_at = timezone.now()
    batch.summary = {**(batch.summary or {}), "created": created, "staged": staged, "skipped": skipped, "failed": failed, "new_terms": created_terms}
    batch.save(update_fields=["status", "completed_at", "summary"])
    audit.log(request or user, "import.completed", entity_type="import", entity_id=str(batch.pk), label=batch.original_name, changes=batch.summary)
    return batch.summary


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------
EXPORT_COLUMNS = [
    ("cast_id", "Cast ID", None),
    ("full_name", "Full name", None),
    ("screen_name", "Screen name", None),
    ("categories", "Categories", None),
    ("tags", "Tags", None),
    ("skills", "Skills", None),
    ("languages", "Languages", None),
    ("gender", "Gender", None),
    ("age", "Age", None),
    ("playing_age", "Playing age", None),
    ("height_cm", "Height (cm)", None),
    ("city", "City", None),
    ("district", "District", None),
    ("state", "State", None),
    ("experience_level", "Experience", None),
    ("travel_willingness", "Travel", None),
    ("instagram", "Instagram", "contact"),
    ("phone", "Phone", "contact"),
    ("alt_phone", "Alternate phone", "contact"),
    ("email", "Email", "contact"),
    ("expected_fee", "Expected fee", "fees"),
    ("fee_type", "Fee type", "fees"),
    ("fee_negotiable", "Negotiable", "fees"),
    ("videos", "Featured videos", None),
    ("completeness", "Profile completeness %", "internal"),
    ("status", "Status", "internal"),
    ("internal_notes", "Internal performance notes", "internal"),
    ("last_note", "Latest contact note", "internal"),
    ("created_at", "Added on", "internal"),
]


def _row_for(t, include):
    d = {
        "cast_id": t.cast_id,
        "full_name": t.full_name,
        "screen_name": t.screen_name,
        "categories": ", ".join(c.name for c in t.categories.all()),
        "tags": ", ".join(c.name for c in t.tags.all()),
        "skills": ", ".join(c.name for c in t.skills.all()),
        "languages": ", ".join(f"{tl.language.name}{' (' + tl.accents + ')' if tl.accents else ''}" for tl in t.languages.all()),
        "gender": t.get_gender_display() if t.gender else "",
        "age": t.age or "",
        "playing_age": f"{t.playing_age_min or ''}–{t.playing_age_max or ''}" if (t.playing_age_min or t.playing_age_max) else "",
        "height_cm": t.height_cm or "",
        "city": t.city,
        "district": t.district,
        "state": t.state,
        "experience_level": t.get_experience_level_display() if t.experience_level else "",
        "travel_willingness": t.get_travel_willingness_display() if t.travel_willingness else "",
        "instagram": t.instagram,
        "phone": t.phone,
        "alt_phone": t.alt_phone,
        "email": t.email,
        "expected_fee": ("Unknown" if t.fee_unknown else (str(t.expected_fee) if t.expected_fee is not None else "")),
        "fee_type": t.get_fee_type_display() if t.fee_type else "",
        "fee_negotiable": {True: "Yes", False: "No", None: ""}[t.fee_negotiable],
        "videos": " | ".join(v.url for v in t.featured_videos.all() if v.url),
        "completeness": t.completeness,
        "status": t.get_status_display(),
        "internal_notes": t.internal_notes,
        "last_note": "",
        "created_at": timezone.localtime(t.created_at).strftime("%Y-%m-%d"),
    }
    if "internal" in include:
        n = t.contact_notes.first()
        if n:
            d["last_note"] = f"{timezone.localtime(n.interaction_at):%d %b %Y} {n.get_kind_display()}: {n.content[:200]}"
    return d


def build_export(job, queryset, include, request):
    """include: set of {'contact', 'fees', 'internal'}"""
    cols = [(k, label) for k, label, group in EXPORT_COLUMNS if group is None or group in include]
    if "client" == job.params.get("audience"):
        queryset = queryset.filter(consent_hold=False)
    qs = queryset.prefetch_related("categories", "tags", "skills", "languages__language", "featured_videos", "contact_notes")
    rows = [_row_for(t, include) for t in qs.iterator(chunk_size=200)]
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    if job.format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([label for _, label in cols])
        for r in rows:
            w.writerow([r[k] for k, _ in cols])
        content = ContentFile(("﻿" + buf.getvalue()).encode("utf-8"))
        name = f"creativheads-talent-{stamp}.csv"
    else:
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Talent"
        ws.append([label for _, label in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F1A17")
        for r in rows:
            ws.append([r[k] for k, _ in cols])
        for i, (k, label) in enumerate(cols, start=1):
            width = min(50, max(len(label), *(len(str(r[k])) for r in rows[:200])) + 2) if rows else len(label) + 2
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
        ws.freeze_panes = "C2"
        buf = io.BytesIO()
        wb.save(buf)
        content = ContentFile(buf.getvalue())
        name = f"creativheads-talent-{stamp}.xlsx"
    job.filename = name
    job.row_count = len(rows)
    job.file.save(name, content, save=False)
    job.save()
    return job


BACKUP_MODELS = [
    "auth.user", "core.systemsetting", "core.sequence", "core.auditlog",
    "talent.term", "talent.talent", "talent.talentlanguage", "talent.talentphoto", "talent.featuredvideo", "talent.portfoliolink",
    "talent.favorite", "talent.contactnote", "talent.guardian", "talent.consent", "talent.retentionreview",
    "casting.project", "casting.role", "casting.shortlistentry", "casting.shortlistevent", "casting.availability", "casting.booking",
]


def build_backup(job, include_media=True):
    """Full structured export: JSON per model + manifest (+ media) in one ZIP.

    Written straight to private export storage so large media sets don't sit in memory.
    """
    from django.apps import apps

    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    job.filename = f"creativheads-full-backup-{stamp}.zip"
    rel = f"{timezone.localtime():%Y/%m}/{job.filename}"
    storage = job.file.storage
    full_path = storage.path(rel)
    os.makedirs(os.path.dirname(full_path), exist_ok=True)
    manifest = {"created_at": timezone.now().isoformat(), "app": "creativheads-casting-manager", "models": {}}
    with zipfile.ZipFile(full_path, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as z:
        for label in BACKUP_MODELS:
            model = apps.get_model(label)
            mgr = getattr(model, "all_objects", model.objects)
            qs = mgr.all().order_by("pk")
            fields = None
            if label == "auth.user":  # never export password hashes
                fields = ["username", "first_name", "last_name", "email", "is_active", "is_superuser", "date_joined", "last_login"]
            z.writestr(f"data/{label}.json", dj_serializers.serialize("json", qs, fields=fields))
            manifest["models"][label] = qs.count()
        if include_media:
            root = str(settings.MEDIA_ROOT)
            count = 0
            for dirpath, _, files in os.walk(root):
                for f in files:
                    fp = os.path.join(dirpath, f)
                    z.write(fp, arcname=os.path.join("media", os.path.relpath(fp, root)))
                    count += 1
            manifest["media_files"] = count
        z.writestr("manifest.json", json.dumps(manifest, indent=2))
    job.file.name = rel
    job.row_count = sum(manifest["models"].values())
    job.save()
    return job
